"""Spreadsheet <-> CSV conventions.

Every workbook becomes a folder with one CSV per sheet:

    data/<area>/<name>/<sheet>.csv            cell values (cached results for formula cells)
    data/<area>/<name>/<sheet>.formulas.csv   only if the sheet has formulas: cell,formula

CSV dialect: UTF-8 with BOM (opens correctly in Excel for Chinese/Cyrillic),
RFC-4180 quoting, CRLF row separators. Trailing empty rows/columns are trimmed.
``data/tables.csv`` records which original workbook (asset key + sha256) each
folder was converted from, so `zontar tables sync` can refresh stale CSVs.
"""
from __future__ import annotations

import csv
import datetime as dt
import re
from pathlib import Path

from openpyxl import load_workbook

from . import layout
from .hashing import sha256_file

ENCODING = "utf-8-sig"
REGISTRY_FIELDS = ["csv_dir", "source_key", "source_sha256", "sheets"]
_UNSAFE = re.compile(r'[\\/:*?"<>|\s]+')


def sheet_filename(title: str) -> str:
    return _UNSAFE.sub("_", title).strip("_") or "sheet"


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (dt.datetime, dt.date, dt.time)):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def _rows(ws) -> list[list[str]]:
    ws.reset_dimensions()
    rows = [[_cell(v) for v in r] for r in ws.iter_rows(values_only=True)]
    while rows and not any(rows[-1]):
        rows.pop()
    width = max((max((i + 1 for i, v in enumerate(r) if v), default=0) for r in rows), default=0)
    return [r[:width] + [""] * (width - len(r)) for r in rows]


def write_csv(path: Path, rows: list[list]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding=ENCODING, newline="") as f:
        csv.writer(f, lineterminator="\r\n").writerows(rows)


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding=ENCODING, newline="") as f:
        return list(csv.DictReader(f))


def xlsx_to_csv(xlsx: Path, out_dir: Path) -> list[str]:
    """Convert every sheet of ``xlsx`` into ``out_dir``. Returns the sheet file stems."""
    out_dir.mkdir(parents=True, exist_ok=True)
    for old in out_dir.glob("*.csv"):
        old.unlink()
    values = load_workbook(xlsx, read_only=True, data_only=True)
    formulas = load_workbook(xlsx, read_only=True, data_only=False)
    stems: list[str] = []
    used: set[str] = set()
    try:
        for ws_v, ws_f in zip(values.worksheets, formulas.worksheets):
            stem = sheet_filename(ws_v.title)
            base, n = stem, 2
            while stem in used:
                stem, n = f"{base}_{n}", n + 1
            used.add(stem)
            stems.append(stem)
            write_csv(out_dir / f"{stem}.csv", _rows(ws_v))
            ws_f.reset_dimensions()
            cells = [[c.coordinate, c.value] for row in ws_f.iter_rows() for c in row
                     if isinstance(getattr(c, "value", None), str) and c.value.startswith("=")]
            if cells:
                write_csv(out_dir / f"{stem}.formulas.csv", [["cell", "formula"], *cells])
    finally:
        values.close()
        formulas.close()
    return stems


# --- registry ------------------------------------------------------------------
def _registry_path() -> Path:
    return layout.data_root() / "tables.csv"


def load_registry() -> dict[str, dict[str, str]]:
    p = _registry_path()
    return {r["csv_dir"]: r for r in read_csv(p)} if p.exists() else {}


def save_registry(reg: dict[str, dict[str, str]]) -> None:
    rows = [REGISTRY_FIELDS] + [[r[k] for k in REGISTRY_FIELDS] for _, r in sorted(reg.items())]
    write_csv(_registry_path(), rows)


def convert(source_key: str, csv_dir: str, force: bool = False) -> bool:
    """Convert asset ``source_key`` into ``data/<csv_dir>`` and register it. Returns True if (re)converted."""
    src = layout.asset(source_key)
    reg = load_registry()
    digest = sha256_file(src)
    prev = reg.get(csv_dir)
    if not force and prev and prev["source_sha256"] == digest and prev["source_key"] == source_key:
        return False
    stems = xlsx_to_csv(src, layout.data_root() / csv_dir)
    reg[csv_dir] = {"csv_dir": csv_dir, "source_key": source_key, "source_sha256": digest, "sheets": ";".join(stems)}
    save_registry(reg)
    return True


def sync_all(force: bool = False) -> list[str]:
    """Re-convert every registered workbook whose source changed."""
    changed = []
    for csv_dir, r in load_registry().items():
        if not layout.asset(r["source_key"]).exists():
            print(f"[skip] source missing locally: {r['source_key']} (run `zontar sync pull`)")
            continue
        if convert(r["source_key"], csv_dir, force):
            changed.append(csv_dir)
    return changed
