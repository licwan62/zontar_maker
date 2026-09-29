from pathlib import Path
from copy import copy
from openpyxl import load_workbook

before_path = Path(r"C:\Users\Lenovo\Desktop\ZONTAR_灰色无耳车罩_新项目提示词交接包_v1.0\Ozon上架链接汇总_写回前备份.xlsx")
after_path = Path(r"C:\Users\Lenovo\Desktop\ozon发货\Ozon上架链接汇总_两店对应完成.xlsx")
before = load_workbook(before_path, data_only=False)
after = load_workbook(after_path, data_only=False)
assert before.sheetnames == after.sheetnames

diffs=[]
for name in before.sheetnames:
    a,b=before[name],after[name]
    for row in a.iter_rows():
        for cell in row:
            if cell.value != b[cell.coordinate].value:
                diffs.append((name,cell.coordinate,cell.value,b[cell.coordinate].value))
assert not diffs, diffs[:10]

dst=after["Ozon上架链接汇总"]
src=after["ozon跨境一店"]
keys=["FordFocus-3L","FordFocus-3M","FordFocus-2M","FordFocus-2L","FordFocus-4M"]
assert [dst.cell(r,15).value for r in range(136,141)] == keys
for i in range(5):
    for c in range(2,26):
        assert dst.cell(136+i,c).value == src.cell(159+i,c).value
    assert dst.cell(136+i,1).value == 135+i

errors=[]
for ws in after.worksheets:
    for row in ws.iter_rows():
        for cell in row:
            if cell.data_type == "e": errors.append((ws.title,cell.coordinate,cell.value))

print({
    "old_cells_unchanged": True,
    "new_range": "Ozon上架链接汇总!A136:Y140",
    "new_keys": keys,
    "store1_filter": after["ozon跨境一店"].auto_filter.ref,
    "freeze_panes": {w.title: str(w.freeze_panes) for w in after.worksheets},
    "summary_table": [(t.name,t.ref) for t in dst._tables.values()],
    "formula_errors": errors,
    "final_size": after_path.stat().st_size,
})
