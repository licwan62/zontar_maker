from pathlib import Path
from openpyxl import load_workbook

official_before=load_workbook(Path(r"C:\Users\Lenovo\Desktop\ZONTAR_灰色无耳车罩_新项目提示词交接包_v1.0\汽车罩官方上架模板_写回前备份.xlsx"),data_only=False)
official_after=load_workbook(Path(r"C:\Users\Lenovo\Desktop\汽车罩官方上架模板.xlsx"),data_only=False)
summary_before=load_workbook(Path(r"C:\Users\Lenovo\Desktop\ZONTAR_灰色无耳车罩_新项目提示词交接包_v1.0\Ozon上架链接汇总_写回前备份.xlsx"),data_only=False)
summary_after=load_workbook(Path(r"C:\Users\Lenovo\Desktop\ozon发货\Ozon上架链接汇总_两店对应完成.xlsx"),data_only=False)

def assert_unchanged(before,after,exceptions):
    assert before.sheetnames==after.sheetnames
    for name in before.sheetnames:
        a,b=before[name],after[name]
        for row in a.iter_rows():
            for cell in row:
                if any(name==sn and r1<=cell.row<=r2 for sn,r1,r2 in exceptions):
                    continue
                assert cell.value==b[cell.coordinate].value,(name,cell.coordinate,cell.value,b[cell.coordinate].value)

assert_unchanged(official_before,official_after,[("模板",10,12)])
assert_unchanged(summary_before,summary_after,[("ozon跨境一店",164,166),("Ozon上架链接汇总",136,138)])

skus=["M","S","YS"]
ots=official_after["模板"]
store=summary_after["ozon跨境一店"]
total=summary_after["Ozon上架链接汇总"]
assert [ots.cell(r,2).value for r in range(10,13)]==[f"ToyotaRAV4-{x}" for x in skus]
assert [store.cell(r,15).value for r in range(164,167)]==[f"ToyotaRAV4-{x}" for x in skus]
assert [total.cell(r,15).value for r in range(136,139)]==[f"ToyotaRAV4-{x}" for x in skus]
assert not any(total.cell(r,15).value in {"FordFocus-3L","FordFocus-3M","FordFocus-2M","FordFocus-2L","FordFocus-4M"} for r in range(136,141))
assert [store.cell(r,13).value for r in range(164,167)]==skus

errors=[]
for wb in (official_after,summary_after):
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.data_type=="e": errors.append((ws.title,c.coordinate,c.value))

print({
    "official_rows":"模板!A10:AX12",
    "summary_store_rows":"ozon跨境一店!A164:AB166",
    "summary_total_rows":"Ozon上架链接汇总!A136:Y138",
    "skus":skus,
    "wrong_ford_total_removed":True,
    "old_cells_unchanged":True,
    "store_filter":summary_after["ozon跨境一店"].auto_filter.ref,
    "formula_errors":errors,
})
