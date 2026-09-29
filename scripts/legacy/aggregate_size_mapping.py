from collections import Counter, defaultdict
from openpyxl import load_workbook

p=r"C:\Users\Lenovo\Desktop\ZONTAR_灰色无耳车罩_新项目提示词交接包_v1.0\俄罗斯带销量数据_0928.xlsx"
wb=load_workbook(p,read_only=True,data_only=True)
ws=wb["RU尺寸组&销量"]
vw=Counter(); uvw=Counter(); blanks=Counter()
for u,v,w in ws.iter_rows(min_row=2,min_col=21,max_col=23,values_only=True):
    u="" if u is None else str(u).strip()
    v="" if v is None else str(v).strip()
    w="" if w is None else str(w).strip()
    blanks[(bool(u),bool(v),bool(w))]+=1
    if v and w: vw[(v,w)]+=1
    if u or v or w: uvw[(u,v,w)]+=1
print("NONBLANK_PATTERNS",blanks)
print("V_TO_W")
byv=defaultdict(list)
for (v,w),n in vw.items(): byv[v].append((w,n))
for v in sorted(byv): print(v,sorted(byv[v],key=lambda x:(-x[1],x[0])))
print("TOP_UVW")
for k,n in uvw.most_common(80): print(k,n)
