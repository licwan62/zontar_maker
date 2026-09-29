from openpyxl import load_workbook
b=load_workbook(r"C:\Users\Lenovo\Desktop\ZONTAR_灰色无耳车罩_新项目提示词交接包_v1.0\Ozon上架链接汇总_取消YS人工指定前备份.xlsx",data_only=False)
a=load_workbook(r"C:\Users\Lenovo\Desktop\ozon发货\Ozon上架链接汇总_两店对应完成.xlsx",data_only=False)
allowed={"ozon跨境一店":{(166,c) for c in (11,12,13,14,28)},"Ozon上架链接汇总":{(138,c) for c in (11,12,13,14)}}
outside=[]; diffs=[]
for name in b.sheetnames:
    for row in b[name].iter_rows():
        for c in row:
            if c.value!=a[name][c.coordinate].value:
                d=(name,c.coordinate,c.value,a[name][c.coordinate].value); diffs.append(d)
                if (c.row,c.column) not in allowed.get(name,set()): outside.append(d)
assert not outside,outside
assert [a["ozon跨境一店"].cell(166,c).value for c in (11,12,13)]==["S","YS-410","YM"]
assert [a["Ozon上架链接汇总"].cell(138,c).value for c in (11,12,13)]==["S","YS-410","YM"]
print({"changed":diffs,"outside_scope":outside,"filter":a["ozon跨境一店"].auto_filter.ref})
