from openpyxl import load_workbook

paths = [
 r"C:\Users\Lenovo\Desktop\汽车罩官方上架模板.xlsx",
 r"C:\Users\Lenovo\Desktop\ozon发货\Ozon上架链接汇总_两店对应完成.xlsx",
]
for p in paths:
    wb=load_workbook(p, data_only=False)
    print("\n",p)
    for ws in wb.worksheets:
        hits=[]
        for row in ws.iter_rows():
            vals=[c.value for c in row]
            if any("focus" in str(v).lower() for v in vals if v is not None):
                hits.append((row[0].row, [str(v)[:100] if v is not None else "" for v in vals[:28]]))
        if hits:
            print(ws.title, len(hits))
            for x in hits[:20]: print(x)
    wb.close()
