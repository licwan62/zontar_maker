from pathlib import Path
from openpyxl import load_workbook
from PIL import Image
import csv

root = Path("Toyota_Land_Cruiser_Tozaroa_最终上架包")
csv_path = root / "04_上架资料" / "Toyota_Land_Cruiser_Tozaroa_上架内容.csv"
with csv_path.open(encoding="utf-8-sig", newline="") as file:
    rows = list(csv.DictReader(file))
assert [row["SKU"] for row in rows] == ["S", "M", "L", "XL"]
assert [row["实际发货尺码"] for row in rows] == ["YM", "YL", "YXL", "YXL"]
assert all(180 <= int(row["标题字符数"]) <= 190 for row in rows)
for row in rows:
    tags = row["标签"].split()
    assert len(tags) == int(row["标签数量"]) == 18
    assert max(map(len, tags)) <= 30
    for field in ("主图", "车型图", "共用SKU图"):
        assert (csv_path.parent / row[field]).resolve().is_file(), (row["SKU"], field, row[field])

pngs = list((root / "01_主图").glob("*.png")) + list((root / "02_车型适配图").glob("*.png")) + list((root / "03_SKU共用图").glob("*.png"))
assert len(pngs) == 9
for path in pngs:
    assert Image.open(path).size == (1086, 1448), (path, Image.open(path).size)
assert len(list((root / "06_共用越野车副图").glob("*.png"))) == 8
assert len(list((root / "07_A+_PC端").glob("*.png"))) == 5
assert len(list((root / "08_A+_移动端").glob("*.png"))) == 5

before = load_workbook("Ozon上架链接汇总_LandCruiser同步前备份.xlsx", data_only=False, read_only=True)
after = load_workbook(r"C:\Users\Lenovo\Desktop\ozon发货\Ozon上架链接汇总_两店对应完成.xlsx", data_only=False, read_only=True)
assert before.sheetnames == after.sheetnames
for sheet_name in before.sheetnames:
    left, right = before[sheet_name], after[sheet_name]
    keep_rows = 166 if sheet_name == "ozon跨境一店" else 138 if sheet_name == "Ozon上架链接汇总" else max(left.max_row, right.max_row)
    max_col = max(left.max_column or 0, right.max_column or 0)
    for r in range(1, keep_rows + 1):
        for c in range(1, max_col + 1):
            assert left.cell(r, c).value == right.cell(r, c).value, (sheet_name, r, c, left.cell(r, c).value, right.cell(r, c).value)

store = after["ozon跨境一店"]
summary = after["Ozon上架链接汇总"]
store_keys = [store.cell(r, 15).value for r in range(167, 171)]
summary_keys = [summary.cell(r, 15).value for r in range(139, 143)]
expected = ["ToyotaLandCruiser-S", "ToyotaLandCruiser-M", "ToyotaLandCruiser-L", "ToyotaLandCruiser-XL"]
assert store_keys == expected and summary_keys == expected
assert [store.cell(r, 13).value for r in range(167, 171)] == ["YM", "YL", "YXL", "YXL"]
before.close(); after.close()

all_text = "\n".join(p.read_text(encoding="utf-8", errors="ignore") for p in root.rglob("*.txt"))
assert "YS" not in all_text
print({"rows": len(rows), "images": len(pngs), "gallery": 8, "pc_a_plus": 5, "mobile_a_plus": 5, "summary_keys": expected})
