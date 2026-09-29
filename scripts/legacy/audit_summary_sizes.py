from collections import Counter, defaultdict
from openpyxl import load_workbook

sales=r"C:\Users\Lenovo\Desktop\ZONTAR_灰色无耳车罩_新项目提示词交接包_v1.0\俄罗斯带销量数据_0928.xlsx"
summary=r"C:\Users\Lenovo\Desktop\ozon发货\Ozon上架链接汇总_两店对应完成.xlsx"

swb=load_workbook(sales,read_only=True,data_only=True)
sws=swb["RU尺寸组&销量"]
counts=Counter()
for v,w in sws.iter_rows(min_row=2,min_col=22,max_col=23,values_only=True):
    v="" if v is None else str(v).strip(); w="" if w is None else str(w).strip()
    if v and w: counts[(v,w)]+=1
byv=defaultdict(list)
for (v,w),n in counts.items(): byv[v].append((w,n))
mapping={v:sorted(items,key=lambda x:(-x[1],x[0]))[0][0] for v,items in byv.items()}

wb=load_workbook(summary,read_only=True,data_only=False)
print("MAPPING",mapping)
for name in ["ozon跨境一店","ozon跨境二店","Ozon上架链接汇总"]:
    ws=wb[name]; diffs=[]; matches=0; unmapped=[]
    for r in range(2,ws.max_row+1):
        oz=ws.cell(r,11).value; ship=ws.cell(r,13).value; item=ws.cell(r,15).value
        oz="" if oz is None else str(oz).strip(); ship="" if ship is None else str(ship).strip()
        if not oz and not ship and not item: continue
        if oz in mapping:
            if ship!=mapping[oz]: diffs.append((r,item,oz,ship,mapping[oz]))
            else: matches+=1
        else: unmapped.append((r,item,oz,ship))
    print("\n",name,"matches",matches,"diffs",len(diffs),"unmapped",len(unmapped))
    print("DIFFS",diffs)
    print("UNMAPPED",unmapped)
