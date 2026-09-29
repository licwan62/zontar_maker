from pathlib import Path
from openpyxl import load_workbook

original = Path(r"C:\Users\Lenovo\Desktop\ozon发货\Ozon上架链接汇总_两店对应完成.xlsx")
candidate = Path(r"C:\Users\Lenovo\Desktop\ZONTAR_灰色无耳车罩_新项目提示词交接包_v1.0\Ozon上架链接汇总_两店对应完成_待写回.xlsx")

wo = load_workbook(original, data_only=False)
wc = load_workbook(candidate, data_only=False)
assert wo.sheetnames == wc.sheetnames, (wo.sheetnames, wc.sheetnames)

allowed_sheet = "Ozon上架链接汇总"
diffs=[]
for s in wo.sheetnames:
    a,b=wo[s],wc[s]
    max_r=max(a.max_row,b.max_row)
    max_c=max(a.max_column,b.max_column)
    for r in range(1,max_r+1):
        for c in range(1,max_c+1):
            av=a.cell(r,c).value
            bv=b.cell(r,c).value
            if av != bv:
                if not (s==allowed_sheet and 136 <= r <= 140):
                    diffs.append((s,r,c,av,bv))

assert not diffs, diffs[:20]

target=wc[allowed_sheet]
source=wc["ozon跨境一店"]
expected_keys=["FordFocus-3L","FordFocus-3M","FordFocus-2M","FordFocus-2L","FordFocus-4M"]
assert [target.cell(r,15).value for r in range(136,141)] == expected_keys
for offset in range(5):
    for c in range(2,26):
        assert target.cell(136+offset,c).value == source.cell(159+offset,c).value, (offset,c)
assert [target.cell(r,1).value for r in range(136,141)] == [135,136,137,138,139]

errors=[]
for ws in wc.worksheets:
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value,str) and cell.value.startswith("#") and cell.data_type=="e":
                errors.append((ws.title,cell.coordinate,cell.value))

print({
    "sheets": wc.sheetnames,
    "unchanged_existing_cells": True,
    "added_rows": "Ozon上架链接汇总!A136:Y140",
    "added_keys": expected_keys,
    "formula_errors": errors,
    "original_size": original.stat().st_size,
    "candidate_size": candidate.stat().st_size,
})

for s in wo.sheetnames:
    a,b=wo[s],wc[s]
    print(s, {
        "merged": (len(a.merged_cells.ranges), len(b.merged_cells.ranges)),
        "validations": (len(a.data_validations.dataValidation), len(b.data_validations.dataValidation)),
        "tables": (list(a.tables.keys()), list(b.tables.keys())),
        "freeze": (a.freeze_panes, b.freeze_panes),
        "auto_filter": (a.auto_filter.ref, b.auto_filter.ref),
        "images": (len(a._images), len(b._images)),
        "charts": (len(a._charts), len(b._charts)),
        "row_dimensions": (len(a.row_dimensions), len(b.row_dimensions)),
        "column_dimensions": (len(a.column_dimensions), len(b.column_dimensions)),
        "conditional_formatting": (len(a.conditional_formatting), len(b.conditional_formatting)),
    })
