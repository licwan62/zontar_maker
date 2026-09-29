from pathlib import Path
from openpyxl import load_workbook

files=[
 Path(r"C:\Users\Lenovo\Desktop\ZONTAR_灰色无耳车罩_新项目提示词交接包_v1.0\俄罗斯带销量数据_0928.xlsx"),
 Path(r"C:\Users\Lenovo\Desktop\ozon发货\Ozon上架链接汇总_两店对应完成.xlsx"),
]
terms=("货号","sku","尺码","发货","商品编号","offer","артикул","размер")
for p in files:
    print("\nFILE",p)
    wb=load_workbook(p,read_only=True,data_only=False)
    print("SHEETS",[(ws.title,ws.max_row,ws.max_column) for ws in wb.worksheets])
    for ws in wb.worksheets:
        hits=[]
        for r in range(1,min(ws.max_row,12)+1):
            for c in range(1,ws.max_column+1):
                v=ws.cell(r,c).value
                if v is not None and any(t in str(v).lower() for t in terms):
                    hits.append((ws.cell(r,c).coordinate,str(v)[:120]))
        if hits: print(ws.title,hits)
    wb.close()
