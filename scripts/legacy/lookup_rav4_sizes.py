from openpyxl import load_workbook
p=r"C:\Users\Lenovo\Desktop\ZONTAR_灰色无耳车罩_新项目提示词交接包_v1.0\俄罗斯带销量数据_0928.xlsx"
w=load_workbook(p,read_only=True,data_only=True)
s=w["RU尺寸组&销量"]
out=[]
for idx,row in enumerate(s.iter_rows(min_row=2,min_col=1,max_col=23,values_only=True),start=2):
    make,model=row[0],row[1]
    if "toyota" in str(make).lower() and "rav" in str(model).lower():
        out.append((idx,make,model,row[2],row[3],row[7],row[20],row[21],row[22]))
print(out)
