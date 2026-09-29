import re
from collections import Counter, defaultdict
from openpyxl import load_workbook

sales=r"C:\Users\Lenovo\Desktop\ZONTAR_灰色无耳车罩_新项目提示词交接包_v1.0\俄罗斯带销量数据_0928.xlsx"
summary=r"C:\Users\Lenovo\Desktop\ozon发货\Ozon上架链接汇总_两店对应完成.xlsx"

def norm(v):
    return re.sub(r"[^0-9a-zа-я]+","",str(v or "").lower().replace("ё","е"))

swb=load_workbook(sales,read_only=True,data_only=True)
sws=swb["RU尺寸组&销量"]
source=[]
for row in sws.iter_rows(min_row=2,min_col=1,max_col=23,values_only=True):
    make,model,u,v,w=row[0],row[1],row[20],row[21],row[22]
    if make and model and v and w:
        source.append((norm(make),norm(model),str(u).strip(),str(v).strip(),str(w).strip(),str(make),str(model)))

wb=load_workbook(summary,read_only=True,data_only=True)
for sheet in ("ozon跨境一店","ozon跨境二店"):
    ws=wb[sheet]; proposals=[]; skipped=Counter()
    for r in range(2,ws.max_row+1):
        brand=ws.cell(r,6).value or ws.cell(r,8).value; c=ws.cell(r,3).value; g=ws.cell(r,7).value; q=ws.cell(r,17).value
        oz=ws.cell(r,11).value; aut=ws.cell(r,12).value; ship=ws.cell(r,13).value; item=ws.cell(r,15).value
        if not item or not brand or not oz: skipped["missing"]+=1; continue
        brand_parts=[norm(x) for x in re.split(r"[\n;,/]+",str(brand)) if norm(x)]
        if len(brand_parts)>3: skipped["universal_brand"]+=1; continue
        model_text=norm(" ".join(str(x or "") for x in (c,g,q)))
        matched=[]
        for mk,md,u,v,w,rawmk,rawmd in source:
            make_ok=any(bp==mk or (len(bp)>=4 and (bp in mk or mk in bp)) for bp in brand_parts)
            model_ok=len(md)>=2 and md in model_text
            if make_ok and model_ok and v==str(oz).strip(): matched.append((u,w,rawmk,rawmd))
        if not matched: skipped["no_match"]+=1; continue
        wc=Counter(x[1] for x in matched); uc=Counter(x[0] for x in matched)
        if len(wc)!=1: skipped["ambiguous_ship"]+=1; continue
        neww=wc.most_common(1)[0][0]; newu=uc.most_common(1)[0][0]
        # Require clear automatic-size evidence when variants exist.
        if len(uc)>1 and uc.most_common(1)[0][1] == uc.most_common(2)[1][1]:
            skipped["ambiguous_auto"]+=1; continue
        if str(aut or "").strip()!=newu or str(ship or "").strip()!=neww:
            proposals.append((r,item,brand,str(oz),aut,newu,ship,neww,len(matched),uc,wc))
    print("\nSHEET",sheet,"PROPOSALS",len(proposals),"SKIPPED",dict(skipped))
    for x in proposals: print(x)
