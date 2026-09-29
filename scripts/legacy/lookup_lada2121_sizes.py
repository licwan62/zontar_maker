from openpyxl import load_workbook
p=r"C:\Users\Lenovo\Desktop\ZONTAR_灰色无耳车罩_新项目提示词交接包_v1.0\俄罗斯带销量数据_0928.xlsx"
w=load_workbook(p,read_only=True,data_only=True)
s=w["RU尺寸组&销量"]
for idx,row in enumerate(s.iter_rows(min_row=2,min_col=1,max_col=23,values_only=True),start=2):
    text=" ".join(str(x or "") for x in row[:9]).lower()
    if any(k in text for k in ("2121","niva","нива","4x4")) and any(k in str(row[0]).lower() for k in ("lada","ваз","vaz")):
        print(idx,row[0],row[1],row[2],row[3],row[7],row[20],row[21],row[22])
