from pathlib import Path
from collections import Counter
from openpyxl import load_workbook

p1 = Path(r"C:\Users\Lenovo\Desktop\汽车罩官方上架模板.xlsx")
p2 = Path(r"C:\Users\Lenovo\Desktop\ozon发货\Ozon上架链接汇总_两店对应完成.xlsx")

def s(v, n=80):
    if v is None: return ""
    x = str(v).replace("\n", " / ")
    return x[:n] + ("..." if len(x) > n else "")

for p in (p1, p2):
    wb = load_workbook(p, read_only=False, data_only=False)
    print("\nFILE", p)
    print("SHEETS", [(w.title, w.max_row, w.max_column) for w in wb.worksheets])
    for w in wb.worksheets:
        hits=[]
        for row in w.iter_rows(min_row=1, max_row=min(10,w.max_row)):
            for c in row:
                if c.value is not None and any(k in str(c.value).lower() for k in ("链接","货号","sku","商品名称","名称")):
                    hits.append((c.coordinate,s(c.value)))
        if hits: print("HEADERS", w.title, hits[:30])
    wb.close()

wa = load_workbook(p1, data_only=False)["模板"]
wb2 = load_workbook(p2, data_only=False)
target = wb2["Ozon上架链接汇总"]

print("\nTEMPLATE PRODUCTS")
for r in range(5, wa.max_row+1):
    vals = [wa.cell(r,c).value for c in (2,3,7,14,15,18,19)]
    if any(v is not None for v in vals): print(r, [s(v,60) for v in vals])

print("\nSUMMARY TAIL / COLUMNS")
print([(target.cell(1,c).coordinate, target.cell(1,c).value) for c in range(1,target.max_column+1)])
for r in range(max(2,target.max_row-20), target.max_row+1):
    print(r, [s(target.cell(r,c).value,60) for c in (1,2,3,4,7,8,11,12,13,15,25)])

# Exact overlaps across likely identifiers.
template_sets = {c: {str(wa.cell(r,c).value).strip() for r in range(5,wa.max_row+1) if wa.cell(r,c).value not in (None,"")} for c in (2,3,7,18,19)}
summary_sets = {c: {str(target.cell(r,c).value).strip() for r in range(2,target.max_row+1) if target.cell(r,c).value not in (None,"")} for c in (2,3,4,7,8,11,12,13,15,25)}
print("\nOVERLAPS")
for ca,sa in template_sets.items():
    for cb,sb in summary_sets.items():
        inter=sa&sb
        if inter: print(ca,cb,len(inter),sorted(inter)[:15])
