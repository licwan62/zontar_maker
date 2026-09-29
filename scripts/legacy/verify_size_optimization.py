from openpyxl import load_workbook

before=load_workbook(r"C:\Users\Lenovo\Desktop\ZONTAR_灰色无耳车罩_新项目提示词交接包_v1.0\Ozon上架链接汇总_尺码优化前备份.xlsx",data_only=False)
after=load_workbook(r"C:\Users\Lenovo\Desktop\ozon发货\Ozon上架链接汇总_两店对应完成.xlsx",data_only=False)
rows1={15,152,153,154,155,156,158,161,162,163,164,165}
rows2={14,15,16,19,20,21,23,24}
allowed={
 "ozon跨境一店":{(r,c) for r in rows1 for c in (12,13,28)},
 "ozon跨境二店":{(r,c) for r in rows2 for c in (12,13,28)},
 "Ozon上架链接汇总":{(r,c) for r in (136,137,138) for c in (12,13)},
}
diffs=[]; outside=[]
for name in before.sheetnames:
    a,b=before[name],after[name]
    for r in range(1,max(a.max_row,b.max_row)+1):
        for c in range(1,max(a.max_column,b.max_column)+1):
            if a.cell(r,c).value!=b.cell(r,c).value:
                diffs.append((name,r,c,a.cell(r,c).value,b.cell(r,c).value))
                if (r,c) not in allowed.get(name,set()): outside.append(diffs[-1])
assert not outside,outside[:20]

s1=after["ozon跨境一店"]
assert (s1["K164"].value,s1["L164"].value,s1["M164"].value)==("M","YL","YL")
assert (s1["K165"].value,s1["L165"].value,s1["M165"].value)==("S","YM+","YM")
assert (s1["K166"].value,s1["L166"].value,s1["M166"].value)==("YS","YS","YS")
errors=[]
for ws in after.worksheets:
    for row in ws.iter_rows():
        for cell in row:
            if cell.data_type=="e": errors.append((ws.title,cell.coordinate,cell.value))
print({
 "changed_cells":len(diffs),
 "outside_scope_changes":len(outside),
 "store1_rows":sorted(rows1),
 "store2_rows":sorted(rows2),
 "rav4_mapping":{"M":"YL","S":"YM","YS":"YS(manual)"},
 "store1_filter":after["ozon跨境一店"].auto_filter.ref,
 "formula_errors":errors,
})
