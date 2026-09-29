from pathlib import Path
from openpyxl import load_workbook

FILES = [
    Path(r"C:\Users\Lenovo\Desktop\汽车罩官方上架模板.xlsx"),
    Path(r"C:\Users\Lenovo\Desktop\ozon发货\Ozon上架链接汇总_两店对应完成.xlsx"),
]

def short(v):
    if v is None:
        return None
    s = str(v).replace("\n", "\\n")
    return s if len(s) <= 120 else s[:117] + "..."

for path in FILES:
    print(f"\n=== FILE {path} ===")
    wb = load_workbook(path, read_only=False, data_only=False)
    for ws in wb.worksheets:
        print(f"\n-- {ws.title} rows={ws.max_row} cols={ws.max_column} --")
        for r in range(1, min(ws.max_row, 15) + 1):
            items = []
            for c in range(1, ws.max_column + 1):
                cell = ws.cell(r, c)
                v = cell.value
                if v is not None:
                    items.append(f"{cell.coordinate}={short(v)!r}")
                if cell.hyperlink:
                    items.append(f"{cell.coordinate}.href={cell.hyperlink.target!r}")
            if items:
                print("; ".join(items[:24]))
        urls = []
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                href = cell.hyperlink.target if cell.hyperlink else None
                if href or (isinstance(v, str) and ("http://" in v or "https://" in v)):
                    urls.append((cell.coordinate, short(v), short(href)))
        print("URL_CELLS", len(urls), urls[:30])
    wb.close()
